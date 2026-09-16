import os
import re
from typing import List


def parse_file(filepath: str) -> str:
    ext = filepath.rsplit(".", 1)[-1].lower()
    handlers = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "md": _parse_text,
        "txt": _parse_text,
        "csv": _parse_text,
        "xlsx": _parse_xlsx,
    }
    handler = handlers.get(ext)
    if not handler:
        raise ValueError(f"不支持的文件类型: {ext}")
    return handler(filepath)


def _parse_pdf(filepath: str) -> str:
    import pdfplumber
    text = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text.append(t)
    return "\n".join(text)


def _parse_docx(filepath: str) -> str:
    from docx import Document
    doc = Document(filepath)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _parse_text(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _parse_xlsx(filepath: str) -> str:
    try:
        import pandas as pd
        dfs = pd.read_excel(filepath, sheet_name=None)
        lines = []
        for sheet_name, df in dfs.items():
            lines.append(f"【工作表：{sheet_name}】")
            for _, row in df.iterrows():
                cells = [str(v) for v in row if str(v) != "nan"]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    except ImportError:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"【工作表：{sheet_name}】")
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)


def chunk_text(text: str, chunk_size: int = 512, overlap: int = 64) -> List[str]:
    if not text.strip():
        return []
    paragraphs = re.split(r"\n\s*\n", text)
    chunks = []
    current = ""
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(current) + len(para) + 1 <= chunk_size:
            current = (current + "\n" + para).strip()
        else:
            if current:
                chunks.append(current)
            if len(para) > chunk_size:
                for i in range(0, len(para), chunk_size - overlap):
                    chunk = para[i:i + chunk_size]
                    if chunk:
                        chunks.append(chunk)
                current = ""
            else:
                current = para
    if current:
        chunks.append(current)
    return chunks
